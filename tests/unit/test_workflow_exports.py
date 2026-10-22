"""Multi-table exports retain values, provenance and safe file semantics."""

import csv
import json

import pytest

from iclotspython.application.multiscale import MultiscaleRequest, run_multiscale
from iclotspython.output import OverwritePolicy
from iclotspython.output import workflows
from tests.unit.test_multiscale_service import images

pytestmark = pytest.mark.unit


def test_workbook_csv_report_values_and_checksums(tmp_path):
    from openpyxl import load_workbook

    result = run_multiscale(MultiscaleRequest(images(), "microchannel"))
    manifest = workflows.export_workflow(result, tmp_path)
    assert manifest.status.value == "success"
    assert len(manifest.files) == 7
    assert all(r.sha256 and r.size_bytes for r in manifest.generated_files)
    workbook = load_workbook(next(tmp_path.glob("*.xlsx")), read_only=True)
    try:
        assert workbook.sheetnames == [t.name for t in result.tables] + ["Reproducibility"]
        assert tuple(workbook["Measurements"].values)[1] == result.table("Measurements").rows[0]
    finally:
        workbook.close()
    with next(tmp_path.glob("*Measurements.csv")).open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[0] == list(result.table("Measurements").columns)
    assert float(rows[1][-1]) == result.table("Measurements").rows[0][-1]
    report = json.loads(next(tmp_path.glob("*.json")).read_text(encoding="utf-8"))
    assert report["provenance"]["scientific_policy"] == "legacy-compatible"
    assert report["parameters"]["thresholds"] == {"red": 50, "green": 50, "blue": 50}
    original = {p.name: p.read_bytes() for p in tmp_path.iterdir()}
    assert workflows.export_workflow(result, tmp_path).status.value == "failed"
    assert original == {p.name: p.read_bytes() for p in tmp_path.iterdir()}


def test_failed_replacement_preserves_original_and_cleans_temporary(tmp_path, monkeypatch):
    result = run_multiscale(MultiscaleRequest(images(), "device"))
    workflows.export_workflow(result, tmp_path, ("csv",))
    original = next(tmp_path.glob("*.csv")).read_bytes()

    def broken(path, table):
        path.write_text("partial", encoding="utf-8")
        raise RuntimeError("writer failed")

    monkeypatch.setattr(workflows, "_write_table", broken)
    manifest = workflows.export_workflow(result, tmp_path, ("csv",), overwrite=OverwritePolicy.REPLACE)
    assert manifest.status.value == "partial"
    assert next(tmp_path.glob("*.csv")).read_bytes() == original
    assert not list(tmp_path.glob(".iclots-*"))
    assert manifest.files[0].issues[0].code == "write_failed"


def test_partial_manifest_and_nonfinite_serialization(tmp_path, monkeypatch):
    from dataclasses import replace

    result = run_multiscale(MultiscaleRequest(images(), "device"))
    result = replace(result, parameters=(("nonfinite", float("inf")),))

    def missing(*args):
        raise ImportError("optional dependency absent")

    monkeypatch.setattr(workflows, "_write_workbook", missing)
    manifest = workflows.export_workflow(result, tmp_path, ("xlsx",))
    assert manifest.status.value == "partial"
    assert manifest.files[0].issues[0].code == "dependency_unavailable"
    assert not list(tmp_path.glob("*.xlsx"))
    assert json.loads(next(tmp_path.glob("*.json")).read_text())["parameters"]["nonfinite"] == "inf"
