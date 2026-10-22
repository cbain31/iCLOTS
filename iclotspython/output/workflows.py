"""Standard multi-table exports, figures and reproducibility reports.

Optional format dependencies load only when executing their writers. Every file
is written to a temporary sibling first; failed replacements preserve originals.
"""

from __future__ import annotations

import csv
from dataclasses import asdict, is_dataclass
import json
import math
import os
from pathlib import Path
import tempfile

from iclotspython.presentation.workflows import WorkflowResult

from .contracts import IssueSeverity, OutputFileRecord, OutputFileStatus, OutputIssue, OverwritePolicy
from .errors import InvalidOutputRequestError
from .paths import build_manifest, prepare_destination, sanitize_stem, success_record


def _json_value(value):
    if is_dataclass(value):
        return _json_value(asdict(value))
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [_json_value(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def result_report(result):
    return _json_value({
        "workflow_id": result.workflow_id, "parameters": dict(result.parameters),
        "provenance": dict(result.provenance), "warnings": result.warnings,
        "tables": [{"name": t.name, "columns": t.data.columns, "row_count": len(t.data.rows)} for t in result.tables],
        "nonfinite_convention": "NaN/Inf preserved as text; never replaced with zero",
    })


def create_workflow_figure(result):
    from matplotlib.figure import Figure

    count = max(len(result.plots), 1)
    figure = Figure(figsize=(6 * min(count, 2), 4 * ((count + 1) // 2)), tight_layout=True)
    for index, panel in enumerate(result.plots):
        axis = figure.add_subplot((count + 1) // 2, min(count, 2), index + 1)
        for series in panel.series:
            if panel.kind == "scatter":
                axis.scatter(series.x, series.y, s=12, label=series.name)
            else:
                axis.plot(series.x, series.y, marker="o", markersize=3, label=series.name)
        axis.set(title=panel.title, xlabel=panel.x_label, ylabel=panel.y_label)
        if panel.series:
            axis.legend()
    if not result.plots:
        axis = figure.add_subplot(111)
        axis.text(0.5, 0.5, "No plot data", ha="center", va="center")
    return figure


def _cell(value):
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(_json_value(value), ensure_ascii=False)


def _write_table(path, table):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(table.columns)
        writer.writerows(tuple(_cell(v) for v in row) for row in table.rows)


def _write_workbook(path, result):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for named in result.tables:
        sheet = workbook.create_sheet(named.name[:31])
        sheet.append(named.data.columns)
        for row in named.data.rows:
            sheet.append(tuple(_cell(v) for v in row))
        sheet.freeze_panes = "A2"
    sheet = workbook.create_sheet("Reproducibility")
    sheet.append(("Field", "Value"))
    for field, value in result_report(result).items():
        sheet.append((field, _cell(value)))
    # Source labels are data, never spreadsheet formulas.
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    try:
        workbook.save(path)
    finally:
        workbook.close()


def _write_plot(path, result):
    from matplotlib.backends.backend_agg import FigureCanvasAgg

    figure = create_workflow_figure(result)
    try:
        FigureCanvasAgg(figure).print_figure(str(path), format="png", dpi=150)
    finally:
        figure.clear()


def _write_report(path, result):
    with path.open("w", encoding="utf-8") as stream:
        json.dump(result_report(result), stream, indent=2, ensure_ascii=False, allow_nan=False)


def output_targets(result, formats):
    stem = sanitize_stem(result.suggested_export_stem, result.workflow_id)
    prefix = f"{stem}_{result.workflow_id}"
    targets = []
    if "csv" in formats:
        targets.extend((f"{prefix}_{sanitize_stem(t.name, 'table')}.csv", "csv", t.data) for t in result.tables)
    if "xlsx" in formats:
        targets.append((f"{prefix}.xlsx", "xlsx", result))
    if "png" in formats:
        targets.append((f"{prefix}.png", "png", result))
    targets.append((f"{prefix}_report.json", "json", result))
    return tuple(targets)


def export_workflow(result: WorkflowResult, destination, formats=("csv", "xlsx", "png"), *, overwrite=OverwritePolicy.FAIL_IF_EXISTS):
    if (not formats or len(set(formats)) != len(formats) or set(formats) - {"csv", "xlsx", "png"}
            or not isinstance(overwrite, OverwritePolicy)):
        raise InvalidOutputRequestError("invalid_formats", "Select unique CSV, XLSX or PNG formats and an explicit overwrite policy.", "formats")
    directory = prepare_destination(destination, False)
    records = []
    writers = {"csv": _write_table, "xlsx": _write_workbook, "png": _write_plot, "json": _write_report}
    for name, format_name, payload in output_targets(result, formats):
        target = directory / name
        temporary = None
        existed = target.exists()
        try:
            if existed and overwrite is OverwritePolicy.FAIL_IF_EXISTS:
                raise FileExistsError(name)
            descriptor, temporary_name = tempfile.mkstemp(dir=directory, prefix=".iclots-", suffix=f".{format_name}")
            os.close(descriptor)
            temporary = Path(temporary_name)
            writers[format_name](temporary, payload)
            if overwrite is OverwritePolicy.REPLACE:
                os.replace(temporary, target)
            elif os.name == "nt":
                # Windows rename refuses existing targets; hard links need
                # privileges unavailable in some desktop sandboxes.
                os.rename(temporary, target)
            else:
                # Atomic publication that refuses a concurrently created target.
                os.link(temporary, target)
            records.append(success_record(kind="report" if format_name == "json" else "plot" if format_name == "png" else "data",
                format_name=format_name, path=target, suggested_filename=name,
                status=OutputFileStatus.REPLACED if existed else OutputFileStatus.CREATED))
        except Exception as exc:
            code = "target_exists" if isinstance(exc, FileExistsError) else "dependency_unavailable" if isinstance(exc, ImportError) else "write_failed"
            issue = OutputIssue(IssueSeverity.ERROR, code, f"Could not export {name}.", f"{type(exc).__name__}: {exc}", str(target), format_name)
            records.append(OutputFileRecord("data", format_name, str(target), name, OutputFileStatus.FAILED, issues=(issue,)))
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
    return build_manifest(result.workflow_id, directory, records)
